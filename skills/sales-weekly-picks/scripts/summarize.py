#!/usr/bin/env python3
"""summarize.py — 고객·상담기록·구매기록 CSV(또는 xlsx)에서 고객별 지표 표를 뽑는다. 표준 라이브러리만 사용.

  python summarize.py --dir . [--today 2026-09-16] [--format md|csv]

xlsx 는 openpyxl 이 설치돼 있을 때만 읽는다(없으면 CSV 사용 안내).
"""
import argparse
import csv
import datetime as dt
import os
import sys

SHEETS = ["고객", "상담기록", "구매기록"]


def parse_date(v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s[:19] if "%H" in fmt else s[:10], fmt).date()
        except ValueError:
            continue
    return None


def load_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return [dict(r) for r in csv.DictReader(f)]


def load_xlsx(path):
    try:
        import openpyxl  # type: ignore
    except ImportError:
        sys.exit("xlsx 를 읽으려면 openpyxl 이 필요합니다. 또는 시트 메뉴에서 CSV 내보내기를 쓰세요.")
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for name in SHEETS:
        if name not in wb.sheetnames:
            out[name] = []
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        if not rows:
            out[name] = []
            continue
        headers = [str(h or "").strip() for h in rows[0]]
        out[name] = [dict(zip(headers, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]
    return out


def load(dir_):
    csvs = {n: os.path.join(dir_, n + ".csv") for n in SHEETS}
    if all(os.path.exists(p) for p in csvs.values()):
        return {n: load_csv(p) for n, p in csvs.items()}
    xlsx = [f for f in os.listdir(dir_) if f.lower().endswith(".xlsx")]
    if xlsx:
        return load_xlsx(os.path.join(dir_, xlsx[0]))
    sys.exit("입력 파일이 없습니다. 시트 메뉴 [고객관리] > [CSV 내보내기] 로 고객.csv / 상담기록.csv / 구매기록.csv 를 만드세요.")


def num(v):
    try:
        return float(str(v).replace(",", "")) if v not in (None, "") else 0.0
    except ValueError:
        return 0.0


def next_anniv(d, today):
    if not d:
        return None
    y = today.year
    try:
        c = d.replace(year=y)
    except ValueError:
        c = dt.date(y, 2, 28)
    if c < today:
        try:
            c = d.replace(year=y + 1)
        except ValueError:
            c = dt.date(y + 1, 2, 28)
    return c


def summarize(data, today):
    rows = []
    for c in data["고객"]:
        cid = str(c.get("고객ID", "")).strip()
        contacts = [x for x in data["상담기록"] if str(x.get("고객ID", "")).strip() == cid]
        purchases = [x for x in data["구매기록"] if str(x.get("고객ID", "")).strip() == cid]
        dates = [parse_date(c.get("마지막연락일"))] + [parse_date(x.get("상담일")) for x in contacts] + [parse_date(x.get("구매일")) for x in purchases]
        dates = [d for d in dates if d]
        last = max(dates) if dates else None
        quote_wait = [(today - parse_date(x.get("상담일"))).days for x in contacts
                      if x.get("유형") == "견적발송" and x.get("결과") == "대기" and parse_date(x.get("상담일"))]
        rebuy = []
        for p in purchases:
            nd = parse_date(p.get("다음구매예정일"))
            if not nd and num(p.get("재구매주기(일)")) and parse_date(p.get("구매일")):
                nd = parse_date(p.get("구매일")) + dt.timedelta(days=int(num(p.get("재구매주기(일)"))))
            if nd:
                rebuy.append((nd - today).days)
        expiry = [(parse_date(p.get("계약만료일/납기일")) - today).days for p in purchases if parse_date(p.get("계약만료일/납기일"))]
        unpaid = [p for p in purchases if str(p.get("수금상태", "")).strip() not in ("완료", "")]
        unpaid_amt = sum(num(p.get("금액")) for p in unpaid)
        unpaid_over = [(today - parse_date(p.get("수금예정일"))).days for p in unpaid if parse_date(p.get("수금예정일"))]
        anniv = [(next_anniv(parse_date(c.get(k)), today) - today).days for k in ("생일", "기념일") if parse_date(c.get(k))]
        total = sum(num(p.get("금액")) for p in purchases)
        recent = sum(num(p.get("금액")) for p in purchases if parse_date(p.get("구매일")) and (today - parse_date(p.get("구매일"))).days <= 183)
        rows.append({
            "고객ID": cid, "고객명": c.get("고객명", ""), "단계": c.get("단계", ""), "의향": c.get("구매의향", ""),
            "마지막연락": last.isoformat() if last else "", "미접촉일": (today - last).days if last else "",
            "견적대기일": max(quote_wait) if quote_wait else "",
            "재구매D": min([d for d in rebuy if d >= 0], default=""),
            "만료D": min([d for d in expiry if d >= 0], default=""),
            "미수금": int(unpaid_amt) if unpaid_amt else "", "미수경과일": max(unpaid_over) if unpaid_over else "",
            "기념일D": min([d for d in anniv if d >= 0], default=""),
            "누적매출": int(total), "최근6개월": int(recent),
            "고민": c.get("고민포인트", ""), "취향": c.get("취향/선호", ""), "다음제안": c.get("다음제안아이디어", ""),
        })
    return rows


def main():
    # Windows 콘솔(cp949)에서도 한글·특수문자가 깨지지 않도록
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=".")
    ap.add_argument("--today", default=dt.date.today().isoformat())
    ap.add_argument("--format", default="md", choices=["md", "csv"])
    a = ap.parse_args()
    today = parse_date(a.today)
    data = load(a.dir)
    rows = summarize(data, today)
    cols = ["고객ID", "고객명", "단계", "의향", "마지막연락", "미접촉일", "견적대기일", "재구매D", "만료D", "미수금", "미수경과일", "기념일D", "누적매출", "최근6개월", "고민", "취향", "다음제안"]
    if a.format == "csv":
        w = csv.DictWriter(sys.stdout, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in cols})
        return
    print(f"기준일 {today} · 고객 {len(rows)} · 상담 {len(data['상담기록'])} · 구매 {len(data['구매기록'])}\n")
    print("| " + " | ".join(cols[:14]) + " |")
    print("|" + "---|" * 14)
    for r in sorted(rows, key=lambda r: (r["미접촉일"] if r["미접촉일"] != "" else -1), reverse=True):
        print("| " + " | ".join(str(r.get(k, "")) for k in cols[:14]) + " |")
    print("\n## 정성 정보")
    for r in rows:
        if r["고민"] or r["취향"] or r["다음제안"]:
            print(f"- **{r['고객명']}** | 고민: {r['고민'] or '-'} / 취향: {r['취향'] or '-'} / 다음 제안: {r['다음제안'] or '-'}")


if __name__ == "__main__":
    main()
